import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
from copy import copy
import tkinter as tk
from tkinter import filedialog
import re
import os
import sys
import json

# ------------------ CACHE ------------------

CACHE_FILE = "ik_cache.json"

def load_ik_cache():
    """Wczytuje zapisane IK z pliku JSON (jeśli istnieje)."""
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_ik_cache(cache):
    """Zapisuje IK do pliku JSON."""
    try:
        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(cache, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"Błąd zapisu cache: {e}")

# ------------------ HELPERY ------------------

def get_resource_path(relative_path):
    """
    Zwraca ścieżkę do zasobu, niezależnie od tego, czy program
    działa w trybie deweloperskim, czy po spakowaniu przez PyInstaller.
    """
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

def copy_cell_format(source_cell, target_cell):
    """
    Kopiuje formatowanie oraz formuły, dynamicznie dostosowując kolumny.
    """
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.fill = copy(source_cell.fill)
        target_cell.number_format = copy(source_cell.number_format)
        target_cell.protection = copy(source_cell.protection)
        target_cell.alignment = copy(source_cell.alignment)
    
    try:
        if source_cell.data_type == 'f':
            old_col_letter = get_column_letter(source_cell.column).upper()
            new_col_letter = get_column_letter(target_cell.column).upper()
            pattern = re.compile(r'(\$?){}'.format(re.escape(old_col_letter)))
            new_formula = pattern.sub(r'\1{}'.format(new_col_letter), source_cell.formula)
            target_cell.formula = new_formula
        else:
            target_cell.value = source_cell.value
    except AttributeError:
        target_cell.value = source_cell.value

# ------------------ POPUP DO IK ------------------

def ask_for_ik_values(articles, parent, cache):
    """
    Tworzy okno Toplevel (popup) z tabelą do wpisania IK.
    Jeśli artykuł jest w cache, IK zostanie wstępnie wpisane.
    """
    ik_values = {}

    def submit():
        for idx, (nr_artykulu, nazwa) in enumerate(articles):
            val = entries[idx].get().strip()
            ik_values[nr_artykulu] = val
            if val:  # zapisujemy tylko niepuste
                cache[str(nr_artykulu)] = val
        save_ik_cache(cache)  # aktualizacja cache po zatwierdzeniu
        dialog.destroy()

    def cancel():
        for nr_artykulu, _ in articles:
            ik_values[nr_artykulu] = cache.get(str(nr_artykulu), "")
        dialog.destroy()

    dialog = tk.Toplevel(parent)
    dialog.title("Podaj IK (opakowanie zbiorcze)")

    tk.Label(dialog, text="Nr artykułu", font=("Arial", 10, "bold"), width=15).grid(row=0, column=0, padx=5, pady=5)
    tk.Label(dialog, text="Nazwa", font=("Arial", 10, "bold"), width=30).grid(row=0, column=1, padx=5, pady=5)
    tk.Label(dialog, text="IK", font=("Arial", 10, "bold"), width=10).grid(row=0, column=2, padx=5, pady=5)

    entries = []
    for i, (nr_artykulu, nazwa) in enumerate(articles, start=1):
        tk.Label(dialog, text=str(nr_artykulu), width=15, anchor="w").grid(row=i, column=0, padx=5, pady=2)
        tk.Label(dialog, text=str(nazwa), width=30, anchor="w").grid(row=i, column=1, padx=5, pady=2)

        entry = tk.Entry(dialog, width=10)
        entry.grid(row=i, column=2, padx=5, pady=2)

        # wstępne wypełnienie z cache
        if str(nr_artykulu) in cache:
            entry.insert(0, cache[str(nr_artykulu)])

        entries.append(entry)

    tk.Button(dialog, text="OK", command=submit, bg="lightgreen", width=15).grid(row=len(articles)+1, column=1, pady=10)
    tk.Button(dialog, text="Anuluj", command=cancel, bg="lightcoral", width=15).grid(row=len(articles)+1, column=2, pady=10)

    parent.wait_window(dialog)
    return ik_values

# ------------------ GŁÓWNA LOGIKA ------------------

def select_file_and_process():
    """
    Otwiera okno dialogowe do wyboru plików i przetwarza dane.
    """
    root = tk.Tk()
    root.withdraw()

    data_file_path = filedialog.askopenfilename(
        title="Wybierz plik z danymi (dane.xlsx)",
        filetypes=[("Excel files", "*.xlsx *.xls")]
    )

    if not data_file_path:
        print("Operacja anulowana przez użytkownika.")
        return

    szablon_file_path = get_resource_path('szablon rozdzielnik.xlsx')
    if not os.path.exists(szablon_file_path):
        print(f"Błąd: Plik '{szablon_file_path}' nie został znaleziony.")
        return

    print(f"Wybrano plik z danymi: {data_file_path}")
    print(f"Ładowanie szablonu z pakietu: {szablon_file_path}")

    try:
        dane_df = pd.read_excel(data_file_path)
        szablon_wb = openpyxl.load_workbook(szablon_file_path)
        szablon_ws = szablon_wb.active
    except Exception as e:
        print(f"Wystąpił błąd podczas ładowania plików: {e}")
        return

    grouped_dane = dane_df.groupby('nr_artykulu')
    start_col_source = 17 # Kolumna Q
    start_col_target = 17 # Kolumna Q

    # przygotuj listę artykułów (nr, nazwa)
    articles = [(nr, group['nazwa'].iloc[0]) for nr, group in grouped_dane]

    # wczytaj cache
    cache = load_ik_cache()

    # popup z IK (z pamięcią)
    ik_values = ask_for_ik_values(articles, root, cache)

    for i, (nr_artykulu, group) in enumerate(grouped_dane):
        nr_artykulu_col = start_col_target + (3 * i)
        stan_col = nr_artykulu_col + 1
        sprzedaz_col = nr_artykulu_col + 2
        
        for row in range(1, szablon_ws.max_row + 1):
            copy_cell_format(szablon_ws.cell(row=row, column=start_col_source), szablon_ws.cell(row=row, column=nr_artykulu_col))
            copy_cell_format(szablon_ws.cell(row=row, column=start_col_source + 1), szablon_ws.cell(row=row, column=stan_col))
            copy_cell_format(szablon_ws.cell(row=row, column=start_col_source + 2), szablon_ws.cell(row=row, column=sprzedaz_col))

        szablon_ws.cell(row=16, column=nr_artykulu_col, value=nr_artykulu)
        szablon_ws.cell(row=16, column=stan_col, value='stan')
        szablon_ws.cell(row=16, column=sprzedaz_col, value='sprzedaz')

        szablon_ws.cell(row=10, column=nr_artykulu_col, value=nr_artykulu)
        szablon_ws.cell(row=11, column=nr_artykulu_col, value=group['nazwa'].iloc[0])

        # dopisz IK w wierszu 14
        szablon_ws.cell(row=14, column=nr_artykulu_col, value=ik_values.get(nr_artykulu, ""))

        for _, row in group.iterrows():
            nr_sklepu = row['nr_sklepu']
            for j in range(17, szablon_ws.max_row + 1):
                if szablon_ws.cell(row=j, column=2).value == nr_sklepu:
                    szablon_ws.cell(row=j, column=stan_col, value=row['stan'])
                    szablon_ws.cell(row=j, column=sprzedaz_col, value=row['sprzedaz'])
                    break

    output_path = filedialog.asksaveasfilename(
        title="Zapisz plik 'rozdzielnik_z_danymi.xlsx'",
        defaultextension=".xlsx",
        filetypes=[("Excel files", "*.xlsx")],
        initialfile="rozdzielnik_z_danymi.xlsx"
    )

    if not output_path:
        print("Operacja zapisu anulowana przez użytkownika.")
        return

    try:
        szablon_wb.save(output_path)
        print(f"Dane zostały pomyślnie przeniesione i zapisane w pliku '{output_path}'.")
    except PermissionError:
        print("Błąd: Nie można zapisać pliku. Upewnij się, że nie jest on otwarty w innym programie.")

# ------------------ START ------------------

if __name__ == "__main__":
    select_file_and_process()
